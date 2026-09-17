"""周报导出：生成 Markdown / HTML 文本。

- Markdown：周整理三段 + 本周实验（标题/状态/目标/结论）+ 下周规划，图片用相对路径引用
- HTML：同上，且图片以 base64 内嵌，单文件自包含，浏览器打开即可见
"""
from __future__ import annotations

import base64
import re
import markdown

from .. import config
from ..models import experiment as exp_model
from ..utils import time_utils


def _status_label(s: str) -> str:
    return exp_model.STATUS_LABELS.get(s, s)


def extract_image_paths(text: str) -> list[str]:
    """提取文本里引用的图片相对路径（images/xxx.jpg）。"""
    return [m for m in re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text or "") if m.startswith("images/")]


def build_markdown(week_start: str, review: dict, stats: dict, experiments: list[dict],
                   plan_items: list[dict], week_label: str) -> str:
    """生成周报 Markdown 文本。"""
    lines = [
        f"# 实验周报 {week_label}",
        "",
        "## 本周统计",
        "",
        f"- 完成实验：{stats['done']}",
        f"- 进行中实验：{stats['running']}",
        f"- 实验记录：{stats['records']} 条",
        f"- 图片：{stats['images']} 张",
        f"- 失败实验：{stats['failed']}",
        f"- 超时实验：{stats['overdue']}",
        "",
        "## 本周总结",
        "",
        review.get("summary") or "（无）",
        "",
        "## 遇到的问题",
        "",
        review.get("problems") or "（无）",
        "",
        "## 下周计划",
        "",
        review.get("next_plan") or "（无）",
        "",
        "## 本周实验",
        "",
    ]
    if experiments:
        for e in experiments:
            lines.append(f"### {e['title']}")
            lines.append("")
            lines.append(f"- 状态：{_status_label(e['status'])}")
            if e.get("planned_start"):
                lines.append(f"- 计划时间：{e['planned_start']}")
            if e.get("type"):
                lines.append(f"- 类型：{e['type']}")
            if e.get("goal"):
                lines.append(f"- 目标：{e['goal']}")
            if e.get("conclusion"):
                lines.append(f"- 结论：{e['conclusion']}")
            lines.append("")
    else:
        lines.append("（无）")
    lines += ["", "## 下周实验规划", ""]
    if plan_items:
        for p in plan_items:
            dur = time_utils.format_duration(p["duration_min"])
            lines.append(f"- {p['planned_date']} {p['planned_start'] or '—'} | {p['title']} | {dur}")
    else:
        lines.append("（无）")
    lines.append("")
    return "\n".join(lines)


_HTML_CSS = """
body { font-family: "Microsoft YaHei UI", "PingFang SC", sans-serif; max-width: 820px;
       margin: 40px auto; padding: 0 24px; color: #1F2937; line-height: 1.7; }
h1 { color: #4F8CFF; border-bottom: 2px solid #4F8CFF; padding-bottom: 8px; }
h2 { color: #1F2937; margin-top: 28px; border-bottom: 1px solid #E5EAF2; padding-bottom: 6px; }
h3 { color: #4F8CFF; margin-top: 20px; }
li { margin: 4px 0; }
img { max-width: 100%; border-radius: 6px; margin: 8px 0; }
"""


def _embed_images(html: str) -> str:
    """把 HTML 里的相对图片路径替换成 base64 data URI，单文件自包含。"""
    def repl(m: re.Match) -> str:
        rel = m.group(1)
        p = config.DATA_DIR / rel
        if p.exists():
            ext = p.suffix.lower().lstrip(".")
            mime = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "gif": "gif",
                    "webp": "webp", "bmp": "bmp"}.get(ext, "jpeg")
            b64 = base64.b64encode(p.read_bytes()).decode()
            return f'src="data:image/{mime};base64,{b64}"'
        return m.group(0)

    return re.sub(r'src="(images/[^"]+)"', repl, html)


def build_html(week_start: str, review: dict, stats: dict, experiments: list[dict],
               plan_items: list[dict], week_label: str) -> str:
    """生成周报 HTML 文本（内联 CSS + 图片 base64 内嵌）。"""
    md = build_markdown(week_start, review, stats, experiments, plan_items, week_label)
    body = markdown.markdown(md, extensions=["tables", "sane_lists"])
    body = _embed_images(body)
    return (
        "<!DOCTYPE html>\n<html lang=\"zh\">\n<head>\n"
        "<meta charset=\"utf-8\">\n"
        f"<title>实验周报 {week_label}</title>\n"
        f"<style>{_HTML_CSS}</style>\n"
        "</head>\n<body>\n" + body + "\n</body>\n</html>\n"
    )


def _docx_text_with_images(doc, text: str) -> None:
    """向 Word 文档添加文本，并插入其中的图片（移除 Markdown 图片语法）。"""
    from docx.shared import Inches

    imgs = extract_image_paths(text or "")
    clean = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text or "").strip()
    if clean:
        doc.add_paragraph(clean)
    for rel in imgs:
        p = config.DATA_DIR / rel
        if p.exists():
            try:
                doc.add_picture(str(p), width=Inches(4))
            except Exception:
                pass


def export_docx(week_start: str, review: dict, stats: dict, experiments: list[dict],
                plan_items: list[dict], week_label: str, dest_path) -> None:
    """生成 Word 周报（.docx），含周整理三段、实验名称+结论、图片。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    # 统一字体：中文宋体、英文/数字 Times New Roman
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.add_heading(f"实验周报 {week_label}", level=0)

    doc.add_heading("本周统计", level=1)
    for label, key in (("完成实验", "done"), ("进行中实验", "running"),
                       ("实验记录", "records"), ("图片", "images"),
                       ("失败实验", "failed"), ("超时实验", "overdue")):
        doc.add_paragraph(f"{label}：{stats[key]}", style="List Bullet")

    doc.add_heading("本周总结", level=1)
    _docx_text_with_images(doc, review.get("summary") or "")
    doc.add_heading("遇到的问题", level=1)
    _docx_text_with_images(doc, review.get("problems") or "")
    doc.add_heading("下周计划", level=1)
    _docx_text_with_images(doc, review.get("next_plan") or "")

    doc.add_heading("本周实验", level=1)
    if experiments:
        for e in experiments:
            doc.add_heading(e["title"], level=2)
            doc.add_paragraph(f"状态：{_status_label(e['status'])}")
            if e.get("planned_start"):
                doc.add_paragraph(f"计划时间：{e['planned_start']}")
            if e.get("goal"):
                doc.add_paragraph(f"目标：{e['goal']}")
            if e.get("conclusion"):
                doc.add_paragraph("结论：")
                _docx_text_with_images(doc, e["conclusion"])
    else:
        doc.add_paragraph("（无）")

    doc.add_heading("下周实验规划", level=1)
    if plan_items:
        for p in plan_items:
            dur = time_utils.format_duration(p["duration_min"])
            doc.add_paragraph(
                f"{p['planned_date']} {p['planned_start'] or '—'} | {p['title']} | {dur}",
                style="List Bullet")
    else:
        doc.add_paragraph("（无）")

    doc.save(str(dest_path))
